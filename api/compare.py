import re
import io
import json
import base64
import os
import sys
import traceback
import math

# Add current directory to sys.path to resolve relative imports robustly
current_dir = os.path.dirname(os.path.abspath(__file__))
if current_dir not in sys.path:
    sys.path.insert(0, current_dir)

from flask import Flask, request, jsonify
from flask_cors import CORS
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# We store any import error that happens during startup
startup_error = None
try:
    from rate_limiter import RateLimiter
    import pdfplumber
    from docx import Document
    from semantic_similarity import get_semantic_similarity_scores
except Exception as e:
    startup_error = {
        "error": "startup_import_error",
        "message": str(e),
        "traceback": traceback.format_exc()
    }

app = Flask(__name__)
CORS(app)

if startup_error is None:
    limiter = RateLimiter()
else:
    limiter = None

def extract_text(file_content, filename):
    """Extracts text from different file types with robust decoding and table support."""
    try:
        if not filename:
            return None
        filename_lower = filename.lower()

        if filename_lower.endswith('.txt'):
            # Cascading decode fallback
            try:
                return file_content.decode('utf-8-sig')
            except UnicodeDecodeError:
                # Only try UTF-16 if it starts with a UTF-16 BOM to avoid false positives
                if file_content.startswith(b'\xff\xfe') or file_content.startswith(b'\xfe\xff'):
                    try:
                        return file_content.decode('utf-16')
                    except UnicodeDecodeError:
                        pass
                return file_content.decode('latin-1')
        elif filename_lower.endswith('.pdf'):
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif filename_lower.endswith('.docx'):
            doc = Document(io.BytesIO(file_content))
            # Extract paragraphs text
            paragraphs_text = [paragraph.text for paragraph in doc.paragraphs]
            # Extract table contents cell-by-cell and paragraph-by-paragraph
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            tables_text.append(paragraph.text)

            # Combine paragraphs and tables text, omitting empty segments
            all_segments = paragraphs_text + tables_text
            return "\n".join(segment for segment in all_segments if segment.strip())
        else:
            return None
    except Exception as e:
        # Propagate the full traceback in a temporary debug exception
        raise ValueError(f"Extraction error for {filename}: {str(e)}\nTraceback:\n{traceback.format_exc()}")

def split_sentences(text):
    """Splits text into sentences using simple regex."""
    # Simple punctuation-based split (. ! ?)
    # NOTE: Known limitation - doesn't handle abbreviations like Mr., Dr., etc.
    sentences = re.split(r'(?<=[.!?])\s+', text)
    return [s.strip() for s in sentences if s.strip()]

def normalize_text(text):
    """Normalizes text by stripping excess whitespace and lowercasing."""
    return re.sub(r'\s+', ' ', text).strip().lower()

def tokenize(text):
    """Tokenize the text into words of length 2 or more, matching (?u)\\b\\w\\w+\\b."""
    return re.findall(r'\b\w\w+\b', text.lower())

def compute_tfidf_similarity(sentences_a_norm, sentences_b_norm):
    """
    Computes custom TF-IDF similarity between sentences in A and sentences in B.
    Matches scikit-learn's default TfidfVectorizer(smooth_idf=True, norm='l2') and cosine_similarity.

    Returns:
        best_matches_a_tfidf (list of floats): Best match scores for sentences in A.
        best_matches_b_tfidf (list of floats): Best match scores for sentences in B.
    """
    all_sentences = sentences_a_norm + sentences_b_norm
    n_docs = len(all_sentences)

    if n_docs == 0:
        return [], []

    # Tokenize all documents
    tokenized_docs = [tokenize(doc) for doc in all_sentences]

    # Identify unique vocabulary terms
    vocab = sorted(list(set(term for doc in tokenized_docs for term in doc)))

    # Compute document frequencies (df)
    df = {term: 0 for term in vocab}
    for doc in tokenized_docs:
        unique_terms = set(doc)
        for term in unique_terms:
            df[term] += 1

    # Compute inverse document frequencies (idf) with smooth_idf=True
    # idf(t) = log((1 + n_docs) / (1 + df(t))) + 1
    idf = {}
    for term in vocab:
        idf[term] = math.log((1 + n_docs) / (1 + df[term])) + 1.0

    # Helper to compute L2-normalized TF-IDF vector for a document
    def get_tfidf_vector(doc_tokens):
        # Term frequencies (tf)
        tf = {}
        for token in doc_tokens:
            tf[token] = tf.get(token, 0) + 1

        # tf-idf values
        vec = {}
        for term, f in tf.items():
            vec[term] = f * idf[term]

        # L2 norm calculation
        square_sum = sum(val * val for val in vec.values())
        norm = math.sqrt(square_sum)

        if norm == 0:
            return {}

        return {term: val / norm for term, val in vec.items()}

    # Get L2-normalized TF-IDF vectors
    vectors_a = [get_tfidf_vector(tokens) for tokens in tokenized_docs[:len(sentences_a_norm)]]
    vectors_b = [get_tfidf_vector(tokens) for tokens in tokenized_docs[len(sentences_a_norm):]]

    # Compute cosine similarity matrix (dot products of normalized vectors)
    # sim_matrix[i][j] = similarity of vector i from A and vector j from B
    sim_matrix = []
    for vec_a in vectors_a:
        row = []
        for vec_b in vectors_b:
            dot = 0.0
            # Iterate over the smaller dict to optimize speed
            if len(vec_a) < len(vec_b):
                for term, val_a in vec_a.items():
                    if term in vec_b:
                        dot += val_a * vec_b[term]
            else:
                for term, val_b in vec_b.items():
                    if term in vec_a:
                        dot += val_b * vec_a[term]
            row.append(dot)
        sim_matrix.append(row)

    # Compute best matches for A
    best_matches_a_tfidf = []
    for i in range(len(sentences_a_norm)):
        best_matches_a_tfidf.append(max(sim_matrix[i]) if sim_matrix[i] else 0.0)

    # Compute best matches for B
    best_matches_b_tfidf = []
    for j in range(len(sentences_b_norm)):
        col_vals = [sim_matrix[i][j] for i in range(len(sentences_a_norm))]
        best_matches_b_tfidf.append(max(col_vals) if col_vals else 0.0)

    return best_matches_a_tfidf, best_matches_b_tfidf

@app.errorhandler(Exception)
def handle_exception(e):
    """Global exception handler to guarantee JSON responses for all errors."""
    traceback.print_exc()
    response = {
        "error": "server_error",
        "message": f"An unexpected server error occurred: {str(e)}",
        "traceback": traceback.format_exc()
    }
    status_code = 500
    if hasattr(e, "code"):
        status_code = e.code
    if hasattr(e, "description"):
        response["message"] = e.description
    return jsonify(response), status_code

@app.route('/api/compare', methods=['POST'])
def compare_documents():
    if startup_error:
        return jsonify({
            "error": "startup_error",
            "message": f"Startup / Import error: {startup_error['message']}",
            "traceback": startup_error['traceback']
        }), 500

    try:
        # Rate Limiting
        forwarded = request.headers.get('x-forwarded-for')
        if forwarded:
            client_ip = forwarded.split(',')[0].strip()
        else:
            client_ip = request.remote_addr or "unknown"

        allowed, remaining, reset_at = limiter.check_and_increment(client_ip)

        if not allowed:
            return jsonify({
                "error": "rate_limit_exceeded",
                "message": "You've used all your free comparisons for today.",
                "reset_at": reset_at,
                "remaining": 0
            }), 429

        data = request.json
        if not data:
            return jsonify({"error": "No data provided"}), 400

        doc_a_raw = data.get('doc_a')
        doc_b_raw = data.get('doc_b')

        # Each doc can be {"type": "text", "content": "..."} or {"type": "file", "content": "base64...", "filename": "..."}

        def get_content(doc_input):
            if not doc_input:
                return ""
            if doc_input.get('type') == 'text':
                return doc_input.get('content', '')
            elif doc_input.get('type') == 'file':
                try:
                    content_b64 = doc_input.get('content', '')
                    filename = doc_input.get('filename', '')
                    file_content = base64.b64decode(content_b64)
                    extracted = extract_text(file_content, filename)
                    if extracted is None:
                        raise ValueError(f"Failed to extract text from {filename}")
                    return extracted
                except Exception as e:
                    raise ValueError(str(e))
            return ""

        try:
            text_a = get_content(doc_a_raw)
            text_b = get_content(doc_b_raw)
        except ValueError as e:
            return jsonify({"error": str(e)}), 400

        if not text_a.strip() or not text_b.strip():
            return jsonify({
                "error": "extraction_failed",
                "message": "Couldn't extract any text from this file. It may be a scanned image or contain no readable text — try a different file or paste the text directly."
            }), 400

        sentences_a_raw = split_sentences(text_a)
        sentences_b_raw = split_sentences(text_b)

        if not sentences_a_raw or not sentences_b_raw:
            return jsonify({
                "error": "extraction_failed",
                "message": "Couldn't extract any text from this file. It may be a scanned image or contain no readable text — try a different file or paste the text directly."
            }), 400

        # Normalize for vectorization
        sentences_a_norm = [normalize_text(s) for s in sentences_a_raw]
        sentences_b_norm = [normalize_text(s) for s in sentences_b_raw]

        # Weights and feature flags
        enable_semantic = os.getenv('ENABLE_SEMANTIC_SIMILARITY', 'true').lower() == 'true'
        gemini_api_key = os.getenv('GEMINI_API_KEY')
        tfidf_weight = float(os.getenv('TFIDF_WEIGHT', 0.4))
        semantic_weight = float(os.getenv('SEMANTIC_WEIGHT', 0.6))

        try:
            # Cosine similarity (TF-IDF) using pure-Python implementation
            best_matches_a_tfidf, best_matches_b_tfidf = compute_tfidf_similarity(
                sentences_a_norm,
                sentences_b_norm
            )

            method = "tfidf_only"
            best_matches_a = best_matches_a_tfidf
            best_matches_b = best_matches_b_tfidf

            # Try Semantic Similarity (Gemini)
            if enable_semantic and gemini_api_key:
                try:
                    best_matches_a_semantic, best_matches_b_semantic = get_semantic_similarity_scores(
                        sentences_a_raw,
                        sentences_b_raw,
                        gemini_api_key
                    )

                    # Blend scores
                    # Formula: blended_score = (tfidf_weight * tfidf_score) + (semantic_weight * semantic_score)
                    best_matches_a = [
                        (tfidf_weight * tf_val) + (semantic_weight * sem_val)
                        for tf_val, sem_val in zip(best_matches_a_tfidf, best_matches_a_semantic)
                    ]
                    best_matches_b = [
                        (tfidf_weight * tf_val) + (semantic_weight * sem_val)
                        for tf_val, sem_val in zip(best_matches_b_tfidf, best_matches_b_semantic)
                    ]
                    method = "blended"
                except Exception as semantic_err:
                    print(f"Fallback to TF-IDF only: Gemini failed - {semantic_err}")
                    method = "tfidf_only"

            # Response construction
            res_sentences_a = []
            for i, text in enumerate(sentences_a_raw):
                res_sentences_a.append({"text": text, "match_score": float(best_matches_a[i])})

            res_sentences_b = []
            for i, text in enumerate(sentences_b_raw):
                res_sentences_b.append({"text": text, "match_score": float(best_matches_b[i])})

            # Unified overall similarity: weighted average of best match scores
            # Formula: (sum(score_i * len_i) for i in A + B) / (sum(len_i) for i in A + B)
            def compute_weighted_avg(sentences, scores):
                total_score = 0
                total_weight = 0
                for s, score in zip(sentences, scores):
                    weight = len(s)
                    total_score += score * weight
                    total_weight += weight
                return total_score, total_weight

            sum_score_a, weight_a = compute_weighted_avg(sentences_a_raw, best_matches_a)
            sum_score_b, weight_b = compute_weighted_avg(sentences_b_raw, best_matches_b)

            overall_similarity = ((sum_score_a + sum_score_b) / (weight_a + weight_b)) * 100

            return jsonify({
                "overall_similarity": round(float(overall_similarity), 2),
                "method": method,
                "sentences_a": res_sentences_a,
                "sentences_b": res_sentences_b,
                "remaining": remaining
            })

        except Exception as e:
            traceback.print_exc()
            return jsonify({
                "error": "similarity_computation_failed",
                "message": f"Similarity computation failed: {str(e)}",
                "traceback": traceback.format_exc()
            }), 500

    except Exception as e:
        traceback.print_exc()
        return jsonify({
            "error": "server_error",
            "message": f"An unexpected error occurred: {str(e)}",
            "traceback": traceback.format_exc()
        }), 500

# For Vercel, we need to export the app or use a handler.
if __name__ == '__main__':
    app.run(port=5000, debug=True)
